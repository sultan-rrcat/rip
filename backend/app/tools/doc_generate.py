"""doc.generate tool — MD + DOCX + PDF from one template model.

Input is a small report model: ``title``, ``sections[{heading, body}]``,
optional ``tables[{headers[], rows[][]}]``. Markdown renders with the stdlib;
DOCX (`python-docx`) and PDF (`reportlab`) are LAZY imports — a deployment
without them fails honest per call instead of breaking tool import.

Binaries travel as base64 in ToolResponse.data with Markdown inline as
`output`; durable delivery arrives with file-based artifacts (Q34).

Effect class: sandboxed (bounded compute producing artifacts).

RIP port: no plugin system — direct Tool subclass (ADR-017).
"""

from __future__ import annotations

import base64
import io
from typing import Any

from app.tools.base import Tool, ToolRequest, ToolResponse


def _need_modules() -> tuple[Any, Any]:
    """Lazily import the DOCX/PDF toolchains (ImportError stays catchable)."""
    try:
        from docx import Document  # type: ignore[import-not-found]
    except ImportError as e:
        raise RuntimeError(f"python-docx is not installed: {e}") from e
    try:
        from reportlab.lib.pagesizes import letter  # type: ignore[import-not-found]
        from reportlab.lib.styles import getSampleStyleSheet  # type: ignore[import-not-found]
        from reportlab.platypus import (  # type: ignore[import-not-found]
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as e:
        raise RuntimeError(f"reportlab is not installed: {e}") from e
    return Document, (
        letter,
        getSampleStyleSheet,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )


def render_markdown(title: str, sections: list[dict], tables: list[dict]) -> str:
    lines = [f"# {title}", ""]
    for section in sections:
        lines += [f"## {section['heading']}", "", section["body"], ""]
    for table in tables:
        lines.append(" | ".join(table["headers"]))
        lines.append(" | ".join("---" for _ in table["headers"]))
        for row in table["rows"]:
            lines.append(" | ".join(str(cell) for cell in row))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_docx(title: str, sections: list[dict], tables: list[dict]) -> bytes:
    Document, _ = _need_modules()
    doc = Document()
    doc.add_heading(title, level=0)
    for section in sections:
        doc.add_heading(section["heading"], level=1)
        doc.add_paragraph(section["body"])
    for table in tables:
        grid = [table["headers"], *table["rows"]]
        doc_table = doc.add_table(rows=len(grid), cols=len(table["headers"]))
        for i, row in enumerate(grid):
            for j, cell in enumerate(row):
                doc_table.cell(i, j).text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(title: str, sections: list[dict], tables: list[dict]) -> bytes:
    from html import escape

    from reportlab.lib import colors  # type: ignore[import-not-found]

    _, mods = _need_modules()
    (
        letter,
        getSampleStyleSheet,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    ) = mods
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for section in sections:
        story += [
            Paragraph(escape(section["heading"]), styles["Heading2"]),
            Paragraph(escape(section["body"]).replace("\n", "<br/>"), styles["Normal"]),
            Spacer(1, 6),
        ]
    for table in tables:
        grid = [
            [escape(str(c)) for c in row] for row in [table["headers"], *table["rows"]]
        ]
        wrapped = [[Paragraph(c, styles["Normal"]) for c in row] for row in grid]
        story.append(
            Table(
                wrapped,
                style=TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ]
                ),
            )
        )
        story.append(Spacer(1, 6))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter).build(story)
    return buf.getvalue()


def _valid_sections(sections: Any) -> list[dict] | None:
    if not isinstance(sections, list) or not sections:
        return None
    clean: list[dict] = []
    for item in sections:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("heading"), str)
            or not isinstance(item.get("body"), str)
        ):
            return None
        clean.append({"heading": item["heading"], "body": item["body"]})
    return clean


def _valid_tables(tables: Any) -> list[dict] | None:
    if tables is None:
        return []
    if not isinstance(tables, list):
        return None
    clean: list[dict] = []
    for item in tables:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("headers"), list)
            or not item["headers"]
            or not isinstance(item.get("rows"), list)
        ):
            return None
        clean.append(
            {"headers": [str(h) for h in item["headers"]], "rows": item["rows"]}
        )
    return clean


class DocGenerateTool(Tool):
    tool_id = "doc.generate"
    name = "Doc Generate"
    description = (
        "Render a titled report as Markdown, DOCX and PDF. "
        "Required input fields: title:string, sections:array of {heading:string, body:string}, "
        "tables: optional array of {headers:[string], rows:[[any]]}. "
        'Example: {"title":"Docker Overview","sections":[{"heading":"Intro","body":"..."}],"tables":[]}'
    )
    input_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "sections": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "body": {"type": "string"},
                    },
                    "required": ["heading", "body"],
                },
            },
            "tables": {"type": "array"},
        },
        "required": ["title", "sections"],
    }
    output_schema = {
        "type": "object",
        "properties": {
            "markdown": {"type": "string"},
            "docx_b64": {"type": "string"},
            "pdf_b64": {"type": "string"},
        },
    }
    effect_class = "sandboxed"  # type: ignore[assignment]
    cost_class = "medium"

    def execute(self, request: ToolRequest) -> ToolResponse:
        title = request.input.get("title")
        if not isinstance(title, str) or not title:
            return ToolResponse(
                tool_id=self.tool_id,
                ok=False,
                output=None,
                error="'title' is required in input",
            )
        sections = _valid_sections(request.input.get("sections"))
        if sections is None:
            return ToolResponse(
                tool_id=self.tool_id,
                ok=False,
                output=None,
                error="'sections' must be a non-empty array of {heading, body}",
            )
        tables = _valid_tables(request.input.get("tables"))
        if tables is None:
            return ToolResponse(
                tool_id=self.tool_id,
                ok=False,
                output=None,
                error="'tables' must be an array of {headers[], rows[][]}",
            )
        markdown = render_markdown(title, sections, tables)
        try:
            docx_bytes = render_docx(title, sections, tables)
            pdf_bytes = render_pdf(title, sections, tables)
        except RuntimeError as e:
            return ToolResponse(
                tool_id=self.tool_id, ok=False, output=None, error=str(e)
            )
        except Exception as e:  # noqa: BLE001 - render failure is a tool failure
            return ToolResponse(
                tool_id=self.tool_id, ok=False, output=None, error=f"render failed: {e}"
            )
        return ToolResponse(
            tool_id=self.tool_id,
            ok=True,
            output=markdown,
            data={
                "markdown": markdown,
                "docx_b64": base64.b64encode(docx_bytes).decode("ascii"),
                "pdf_b64": base64.b64encode(pdf_bytes).decode("ascii"),
            },
        )
